"""
可视化模块
功能：生成柱状图、饼图、折线图、堆叠柱状图、组合图，并将图片嵌入 Excel 文件

设计原则：
  - 所有图表函数返回图片路径字符串，便于 reporter 模块串联
  - 自动创建输出目录，调用方无需关心路径是否存在
  - 中文字体自动检测与注册，兼容 Windows / macOS / Linux
  - 图表统一使用专业配色，数值标签自动标注
"""
import matplotlib

matplotlib.use('Agg')

import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import pandas as pd
import numpy as np
from pathlib import Path


def _setup_chinese_font():
    """
    自动检测并注册系统中文字体，解决 matplotlib 中文乱码问题。

    优先顺序：
      1. 手动注册已知的中文字体文件路径（Linux 常见位置）
      2. 从 matplotlib 已识别字体中查找含中文关键字的字体
      3. 兜底使用 DejaVu Sans（英文正常，中文会显示方框）

    该函数在模块加载时自动执行一次。
    """
    # 已知中文字体文件路径（按优先级排列）
    candidate_font_paths = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",  # Linux: Noto Sans CJK
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",  # Linux: 文泉驿微米黑
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",  # Linux: 文泉驿正黑
        "/System/Library/Fonts/PingFang.ttc",  # macOS: 苹方
        "/System/Library/Fonts/STHeiti Light.ttc",  # macOS: 黑体
        "C:/Windows/Fonts/msyh.ttc",  # Windows: 微软雅黑
        "C:/Windows/Fonts/simhei.ttf",  # Windows: 黑体
    ]

    registered = False
    for font_path in candidate_font_paths:
        if Path(font_path).exists():
            try:
                # addfont 会将字体注册到 matplotlib 的字体管理器
                fm.fontManager.addfont(font_path)
                # 从字体文件路径中提取字体名称
                font_prop = fm.FontProperties(fname=font_path)
                font_name = font_prop.get_name()
                # 将该字体设为 sans-serif 首选
                plt.rcParams["font.sans-serif"] = [font_name, "DejaVu Sans", "Arial Unicode MS"]
                registered = True
                print(f"[图表] 中文字体已注册: {font_name} ({font_path})")
                break
            except Exception as e:
                print(f"[图表] 字体注册失败 {font_path}: {e}")
                continue

    if not registered:
        # 从已识别字体中查找
        available = [f.name for f in fm.fontManager.ttflist
                     if any(k in f.name for k in
                            ["Hei", "Song", "Kai", "Ming", "CJK", "Noto Sans CJK", "WenQuanYi", "SimHei",
                             "Microsoft YaHei"])]
        if available:
            plt.rcParams["font.sans-serif"] = [available[0], "DejaVu Sans"]
            print(f"[图表] 使用已识别中文字体: {available[0]}")
        else:
            print("[图表] 警告：未找到中文字体，中文可能显示为方框")

    # 解决负号显示为方块的问题
    plt.rcParams["axes.unicode_minus"] = False
    # 统一图表 DPI
    plt.rcParams["figure.dpi"] = 150


# 模块加载时自动执行字体配置
_setup_chinese_font()

# 主色：深蓝；辅助色：橙、绿、红、紫、棕、粉、灰
_COLORS = ["#4C72B0", "#DD8452", "#55A868", "#C44E52", "#8172B3", "#937860", "#DA8BC3", "#8C8C8C"]


# 1---------柱状图
def bar_chart(
        df: pd.DataFrame,
        x_col: str,
        y_col: str,
        title: str,
        output_path: str,
        top_n: int = 15,
        horizontal: bool = False,
) -> str:
    """
        生成柱状图并保存为 PNG。

        :param df: 输入 DataFrame
        :param x_col: X 轴列名（分类）
        :param y_col: Y 轴列名（数值）
        :param title: 图表标题
        :param output_path: 图片保存路径
        :param top_n: 仅显示前 N 条数据，避免类目过多挤在一起
        :param horizontal: True=水平柱状图（类目名较长时推荐），False=垂直柱状图
        :return: 图片保存路径
        """
    data = df.head(top_n).copy()
    fig, ax = plt.subplots(figsize=(10, 6) if not horizontal else (10, max(6, len(data) * 0.5)))

    if horizontal:
        # 水平柱状图
        bars = ax.barh(data[x_col].astype(str), data[y_col], color=_COLORS[0])
        ax.set_xlabel(y_col)
        ax.set_ylabel(x_col)
        for bar in bars:
            width = bar.get_width()
            ax.text(width, bar.get_y() + bar.get_height() / 2, f"{width:,.0f}", ha="left", va="center", fontsize=9)
        ax.invert_yaxis()
    else:
        # 垂直柱状图
        bars = ax.bar(data[x_col].astype(str), data[y_col], color=_COLORS[0])
        ax.set_xlabel(x_col)
        ax.set_ylabel(y_col)
        # 数值标签放在柱子顶部
        for bar in bars:
            height = bar.get_height()
            ax.text(bar.get_x() + bar.get_width() / 2, height,
                    f"{height:,.0f}", ha="center", va="bottom", fontsize=9)
        plt.xticks(rotation=45, ha="right")

    ax.set_title(title, fontsize=14, fontweight="bold", pad=15)
    ax.grid(axis="y" if not horizontal else "x", alpha=0.3, linestyle="--")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)  # 隐藏上，右边框

    plt.tight_layout()
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()

    print(f"[图表] 柱状图已保存: {output_path}")
    return output_path


# 2---------饼图
def pie_chart(
        df: pd.DataFrame,
        label_col: str,
        value_col: str,
        title: str,
        output_path: str,
        top_n: int = 8
) -> str:
    """
        生成饼图并保存为 PNG。

        饼图适合展示占比结构，但类目不宜过多（建议 ≤8），
        超过 top_n 的部分会被截断。

        :param df: 输入 DataFrame
        :param label_col: 标签列名
        :param value_col: 数值列名
        :param title: 图表标题
        :param output_path: 图片保存路径
        :param top_n: 最多显示 N 个类目
        :return: 图片保存路径
        """
    data = df.head(top_n).copy()
    fig, ax = plt.subplots(figsize=(8, 8))

    max_idx = data[value_col].idxmax()
    explode = [0.05 if i == max_idx else 0 for i in data.index]

    wedges, texts, autotexts = ax.pie(
        data[value_col],
        labels=data[label_col].astype(str),
        autopct="%1.1f%%",  # 百分比格式
        startangle=90,  # 从正上方开始
        colors=_COLORS[:len(data)],  # 使用统一配色
        explode=explode,  # 突出最大块
        pctdistance=0.75,  # 百分比文字距圆心距离
        textprops={"fontsize": 10}
    )

    for autotext in autotexts:
        autotext.set_color("white")
        autotext.set_fontweight("bold")

    ax.set_title(title, fontsize=14, fontweight="bold", pad=15)

    plt.tight_layout()
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()

    print(f"[图表] 饼图已保存: {output_path}")
    return output_path


# 3----------折线图
def line_chart(
        df: pd.DataFrame,
        x_col: str,
        y_col: str,
        title: str,
        output_path: str,
        group_col: str = None
) -> str:
    """
        生成折线图并保存为 PNG。

        :param df: 输入 DataFrame
        :param x_col: X 轴列名（通常是时间/月份）
        :param y_col: Y 轴列名（数值）
        :param title: 图表标题
        :param output_path: 图片保存路径
        :param group_col: 分组列，指定后每条线代表一个组（多系列折线图）
        :return: 图片保存路径
        """
    fig, ax = plt.subplots(figsize=(10, 6))
    if group_col and group_col in df.columns:
        groups = df[group_col].unique()
        for i, group in enumerate(groups):
            subset = df[df[group_col] == group].sort_values(x_col)
            ax.plot(subset[x_col].astype(str), subset[y_col], marker="o", linewidth=2, label=str(group),
                    color=_COLORS[i % len(_COLORS)])
        ax.legend(title=group_col, loc="best", fontsize=9)
    else:
        data = df.sort_values(x_col)
        ax.plot(data[x_col].astype(str), data[y_col],
                marker="o", color=_COLORS[1], linewidth=2)
        # 数据点上方标注数值
        for x, y in zip(data[x_col].astype(str), data[y_col]):
            ax.text(x, y, f"{y:,.0f}", ha="center", va="bottom", fontsize=9)

    ax.set_title(title, fontsize=14, fontweight="bold", pad=15)
    ax.set_xlabel(x_col)
    ax.set_ylabel(y_col)
    plt.xticks(rotation=45, ha="right")
    ax.grid(True, alpha=0.3, linestyle="--")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    plt.tight_layout()
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()

    print(f"[图表] 折线图已保存: {output_path}")
    return output_path


# 4---------堆叠柱状图
def stacked_bar_chart(
        df: pd.DataFrame,
        x_col: str,
        y_col: str,
        stack_col: str,
        title: str,
        output_path: str
) -> str:
    """
        生成堆叠柱状图，展示每个 X 类目下不同子类的构成。

        :param df: 输入 DataFrame（需已聚合，每行一个 x+stack 组合）
        :param x_col: X 轴列名（如月份）
        :param y_col: Y 轴数值列名（如销售金额）
        :param stack_col: 堆叠维度列名（如产品类别）
        :param title: 图表标题
        :param output_path: 图片保存路径
        :return: 图片保存路径
        """
    pivot = df.pivot_table(index=x_col, columns=stack_col, values=y_col, aggfunc="sum").fillna(0)

    fig, ax = plt.subplots(figsize=(10, 6))

    # bottom追踪当前堆叠高度

    bottom = np.zeros(len(pivot))
    x_pos = np.arange(len(pivot.index))

    for i, col in enumerate(pivot.columns):
        values = pivot[col].values
        ax.bar(x_pos, values, bottom=bottom, label=str(col), color=_COLORS[i % len(_COLORS)], edgecolor="white",
               linewidth=0.5)
        bottom += values

    ax.set_xticks(x_pos)
    ax.set_xticklabels(pivot.index.astype(str), rotation=45, ha="right")
    ax.set_title(title, fontsize=14, fontweight="bold", pad=15)
    ax.set_xlabel(x_col)
    ax.set_ylabel(y_col)
    ax.legend(title=stack_col, loc="best", fontsize=9)
    ax.grid(axis="y", alpha=0.3, linestyle="--")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    plt.tight_layout()
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()

    print(f"[图表] 堆叠柱状图已保存: {output_path}")
    return output_path


# 5-----------柱状+折线图
def combo_chart(
        df: pd.DataFrame,
        x_col: str,
        bar_col: str,
        line_col: str,
        title: str,
        output_path: str
) -> str:
    """
    生成柱状+折线组合图，双 Y 轴。

    典型场景：柱状图展示销售额（左轴），折线图展示环比增长率（右轴）。

    :param df: 输入 DataFrame
    :param x_col: X 轴列名
    :param bar_col: 柱状图数值列（左 Y 轴）
    :param line_col: 折线图数值列（右 Y 轴）
    :param title: 图表标题
    :param output_path: 图片保存路径
    :return: 图片保存路径
    """
    data = df.sort_values(x_col)

    fig, ax1 = plt.subplots(figsize=(10, 6))

    # 左轴：柱状图
    bars = ax1.bar(data[x_col].astype(str), data[bar_col],
                   color=_COLORS[0], alpha=0.7, label=bar_col, width=0.6)
    ax1.set_xlabel(x_col)
    ax1.set_ylabel(bar_col, color=_COLORS[0])
    ax1.tick_params(axis="y", labelcolor=_COLORS[0])
    # 柱顶数值
    for bar in bars:
        height = bar.get_height()
        ax1.text(bar.get_x() + bar.get_width() / 2, height,
                 f"{height:,.0f}", ha="center", va="bottom", fontsize=9, color=_COLORS[0])

    # 右轴：折线图
    ax2 = ax1.twinx()
    ax2.plot(data[x_col].astype(str), data[line_col],
             marker="o", color=_COLORS[1], linewidth=2, label=line_col)
    ax2.set_ylabel(line_col, color=_COLORS[1])
    ax2.tick_params(axis="y", labelcolor=_COLORS[1])
    # 折线点数值
    for x, y in zip(data[x_col].astype(str), data[line_col]):
        if pd.notna(y):
            ax2.text(x, y, f"{y:.1f}%", ha="center", va="bottom",
                     fontsize=9, color=_COLORS[1])

    ax1.set_title(title, fontsize=14, fontweight="bold", pad=15)
    plt.xticks(rotation=45, ha="right")
    ax1.spines["top"].set_visible(False)
    ax2.spines["top"].set_visible(False)

    # 合并两个轴的图例
    lines1, labels1 = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines1 + lines2, labels1 + labels2, loc="upper left", fontsize=9)

    plt.tight_layout()
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()

    print(f"[图表] 组合图已保存: {output_path}")
    return output_path


# 6----------图片嵌入Excel
def insert_image_to_excel(
        excel_path: str,
        image_path: str,
        sheet_name: str,
        cell: str = "A1",
        width: int = 480
) -> None:
    """
        将图片插入到指定 Excel 工作表的指定单元格。

        图片按原始宽高比等比例缩放，默认宽度 480px。

        :param excel_path: Excel 文件路径
        :param image_path: 图片文件路径
        :param sheet_name: 目标工作表名称（不存在则自动创建）
        :param cell: 插入位置的单元格，如 "F2"
        :param width: 图片显示宽度（像素），高度按原始比例自动计算
        """
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage
    wb = load_workbook(excel_path)
    if sheet_name not in wb.sheetnames:
        wb.create_sheet(sheet_name)
    ws = wb[sheet_name]

    img = XLImage(image_path)

    try:
        with PILImage.open(image_path) as pil_img:
            orig_w, orig_h = pil_img.size
            if orig_w > 0:
                img.width = width
                img.height = int(width * orig_h / orig_w)
    except Exception:
        # 读取失败时使用默认尺寸
        img.width = width
        img.height = int(width * 0.67)

    ws.add_image(img, cell)
    wb.save(excel_path)

    print(f"[图表] 图片已插入 {excel_path} -> {sheet_name}!{cell} (宽度: {img.width}px)")
